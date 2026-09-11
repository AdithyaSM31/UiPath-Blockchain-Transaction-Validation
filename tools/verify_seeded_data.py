"""
verify_seeded_data.py
=====================
Independent check that the generated dataset really contains the anomalies it claims,
and that the ABI-encoded `input` field decodes cleanly.

This is deliberately a *separate* implementation from the generator: it re-derives the
expected outcome for every transaction from the raw JSON and the ERP workbook, exactly
as the UiPath validation engine will. It therefore doubles as the reference oracle that
Stage 3 results get compared against - if the bot and this script disagree, one of them
is wrong and the run is not signed off.

Run:  python tools/verify_seeded_data.py
Exit code 0 = every expectation met.
"""

from __future__ import annotations

import json
import os
import sys
from collections import defaultdict
from datetime import datetime

from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PROJ = os.path.join(ROOT, "BlockchainLogisticsValidator")
DATA = os.path.join(PROJ, "Data")
INPUT = os.path.join(DATA, "Input")

TOL_PASS_H = 2
TOL_WARN_H = 4

failures: list[str] = []


def check(label: str, condition: bool, detail: str = "") -> None:
    if condition:
        print(f"  PASS  {label}")
    else:
        print(f"  FAIL  {label}   {detail}")
        failures.append(label)


# --------------------------------------------------------------------------
# ABI decoding - the exact operation the UiPath preprocessing step performs
# --------------------------------------------------------------------------
def decode_input(input_hex: str) -> dict:
    """Decode `f(string, uint256, string)` call data into its three arguments."""
    body = bytes.fromhex(input_hex[10:])          # strip '0x' + 4-byte selector
    words = [body[i:i + 32] for i in range(0, len(body), 32)]

    def word_int(i: int) -> int:
        return int.from_bytes(words[i], "big")

    def string_at(byte_offset: int) -> str:
        idx = byte_offset // 32
        length = int.from_bytes(words[idx], "big")
        raw = b"".join(words[idx + 1:])[:length]
        return raw.decode("utf-8")

    return {
        "shipmentId": string_at(word_int(0)),
        "quantity": word_int(1),
        "poNumber": string_at(word_int(2)),
        "selector": input_hex[:10],
    }


# --------------------------------------------------------------------------
# Load
# --------------------------------------------------------------------------
def load_sheet(path: str, sheet: str) -> list[dict]:
    ws = load_workbook(path, data_only=True)[sheet]
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(h) for h in rows[0]]
    return [dict(zip(headers, r)) for r in rows[1:] if r[0] is not None]


def main() -> int:
    chain_path = os.path.join(INPUT, "MockChain", "etherscan_txlist_response.json")
    with open(chain_path, encoding="utf-8") as f:
        payload = json.load(f)

    check("Etherscan envelope has status=1 and a result array",
          payload.get("status") == "1" and isinstance(payload.get("result"), list))

    txs = payload["result"]
    config_path = os.path.join(DATA, "Config.xlsx")
    logistics_path = os.path.join(INPUT, "LogisticsRecords.xlsx")

    wallets = load_sheet(config_path, "ApprovedWallets")
    approved = {str(w["Address"]).lower() for w in wallets if str(w["Active"]).upper() == "TRUE"}
    seq_rows = load_sheet(config_path, "EventSequence")
    selector_to_event = {str(s["MethodId"]).lower(): s["EventName"] for s in seq_rows}
    event_order = {s["EventName"]: int(s["StepOrder"]) for s in seq_rows}

    shipments = {s["ShipmentID"]: s for s in load_sheet(logistics_path, "Shipments")}
    erp_events = load_sheet(logistics_path, "ShipmentEvents")
    erp_lookup = {(e["ShipmentID"], e["EventName"]): e for e in erp_events}

    print("\n--- Structural checks ---")
    check("48 transactions on the mock feed", len(txs) == 48, f"got {len(txs)}")
    check("48 expected ERP milestone lines", len(erp_events) == 48, f"got {len(erp_events)}")
    check("12 shipments in the ERP master", len(shipments) == 12, f"got {len(shipments)}")
    check("6 approved partner wallets", len(approved) == 6, f"got {len(approved)}")

    # --- decode every transaction ----------------------------------------
    decoded = []
    decode_errors = []
    for t in txs:
        try:
            d = decode_input(t["input"])
        except Exception as exc:                       # noqa: BLE001
            decode_errors.append(f"{t['hash'][:12]}: {exc}")
            continue
        d.update(
            hash=t["hash"].lower(),
            sender=t["from"].lower(),
            ts=datetime.utcfromtimestamp(int(t["timeStamp"])),
            event=selector_to_event.get(t["methodId"].lower(), "UNKNOWN"),
        )
        decoded.append(d)

    check("every transaction's ABI input decodes", not decode_errors, "; ".join(decode_errors[:3]))
    check("every selector maps to a known milestone",
          all(d["event"] != "UNKNOWN" for d in decoded))
    check("decoded ShipmentIDs all exist in the ERP master",
          all(d["shipmentId"] in shipments for d in decoded))
    check("decoded PO numbers match the ERP master",
          all(d["poNumber"] == shipments[d["shipmentId"]]["PONumber"] for d in decoded))

    # ----------------------------------------------------------------------
    print("\n--- R1 timestamp drift ---")
    drift = {}
    for d in decoded:
        erp = erp_lookup.get((d["shipmentId"], d["event"]))
        if erp:
            gap = abs((d["ts"] - erp["ExpectedEventTimestampUtc"]).total_seconds()) / 3600.0
            drift[(d["shipmentId"], d["event"])] = gap

    warn = {k for k, g in drift.items() if TOL_PASS_H < g <= TOL_WARN_H}
    fail = {k for k, g in drift.items() if g > TOL_WARN_H}
    check("SHP-1002/InTransit lands in the WARNING band",
          ("SHP-1002", "InTransit") in warn,
          f"drift={drift.get(('SHP-1002','InTransit')):.2f}h")
    check("SHP-1003/InTransit breaches the FAIL threshold",
          ("SHP-1003", "InTransit") in fail,
          f"drift={drift.get(('SHP-1003','InTransit')):.2f}h")
    check("no other transaction drifts past tolerance",
          warn | fail == {("SHP-1002", "InTransit"), ("SHP-1003", "InTransit")},
          f"unexpected: {sorted((warn | fail) - {('SHP-1002','InTransit'), ('SHP-1003','InTransit')})}")

    # ----------------------------------------------------------------------
    print("\n--- R2 quantity mismatch ---")
    mismatched = {(d["shipmentId"], d["event"]): (d["quantity"], shipments[d["shipmentId"]]["ExpectedQty"])
                  for d in decoded
                  if d["quantity"] != shipments[d["shipmentId"]]["ExpectedQty"]}
    check("SHP-1005/Dispatched is short by 20 units",
          mismatched.get(("SHP-1005", "Dispatched")) == (480, 500),
          str(mismatched.get(("SHP-1005", "Dispatched"))))
    check("no other quantity mismatch exists", len(mismatched) == 1, str(mismatched))

    # ----------------------------------------------------------------------
    print("\n--- R3 wallet whitelist ---")
    rogue = {(d["shipmentId"], d["event"], d["sender"]) for d in decoded if d["sender"] not in approved}
    check("exactly one transaction comes from an unapproved wallet", len(rogue) == 1, str(rogue))
    check("it is SHP-1007/Delivered",
          any(s == "SHP-1007" and e == "Delivered" for s, e, _ in rogue), str(rogue))

    # ----------------------------------------------------------------------
    print("\n--- R4 duplicate detection ---")
    by_hash = defaultdict(list)
    for d in decoded:
        by_hash[(d["shipmentId"], d["hash"])].append(d)
    dupes = {k: v for k, v in by_hash.items() if len(v) > 1}
    check("exactly one transaction hash is recorded twice", len(dupes) == 1, str(list(dupes)[:2]))
    check("the duplicate belongs to SHP-1009/Delivered",
          all(k[0] == "SHP-1009" and v[0]["event"] == "Delivered" for k, v in dupes.items()))

    # ----------------------------------------------------------------------
    print("\n--- R5 milestone ordering ---")
    out_of_order = []
    per_ship = defaultdict(list)
    for d in decoded:
        per_ship[d["shipmentId"]].append(d)
    for sid, evs in per_ship.items():
        chron = sorted({(e["event"], e["ts"]) for e in evs}, key=lambda x: x[1])
        steps = [event_order[e] for e, _ in chron]
        if steps != sorted(steps):
            out_of_order.append(sid)
    check("exactly one shipment has milestones out of order",
          out_of_order == ["SHP-1011"], str(out_of_order))

    # ----------------------------------------------------------------------
    print("\n--- R6 required smart-contract event ---")
    missing = [sid for sid in shipments
               if "GoodsReceived" not in {d["event"] for d in per_ship.get(sid, [])}]
    check("exactly one shipment never emits GoodsReceived",
          missing == ["SHP-1012"], str(missing))
    check("its ERP record still expects that milestone",
          ("SHP-1012", "GoodsReceived") in erp_lookup)

    # ----------------------------------------------------------------------
    print("\n--- Clean control group ---")
    flagged = {"SHP-1002", "SHP-1003", "SHP-1005", "SHP-1007", "SHP-1009", "SHP-1011", "SHP-1012"}
    clean = sorted(set(shipments) - flagged)
    check("5 shipments are entirely clean", len(clean) == 5, str(clean))
    print(f"        clean: {', '.join(clean)}")

    # ----------------------------------------------------------------------
    print("\n--- Address normalisation is actually needed ---")
    mixed_case = [t["from"] for t in txs if t["from"] != t["from"].lower()]
    erp_lower = [str(e["HandlerAddress"]) for e in erp_events]
    check("chain addresses are checksummed (mixed case)", len(mixed_case) > 0)
    check("ERP addresses are lowercase",
          all(a == a.lower() for a in erp_lower))
    check("so a raw string compare would fail without normalisation",
          not set(mixed_case) & set(erp_lower))

    print()
    if failures:
        print(f"{len(failures)} CHECK(S) FAILED")
        return 1
    print("All checks passed - dataset is fit for Stage 3 validation.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
